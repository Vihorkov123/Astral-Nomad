# -*- coding: utf-8 -*-
"""Презентация проекта «Астральный скиталец» (.pptx), 16:9, тёмная тема."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

BG     = RGBColor(0x0B, 0x10, 0x24)   # тёмно-синий космос
CARD   = RGBColor(0x16, 0x1D, 0x3D)
EDGE   = RGBColor(0x5A, 0x6A, 0xA8)
TEXT   = RGBColor(0xDF, 0xE6, 0xFF)
MUTED  = RGBColor(0x9A, 0xA6, 0xD0)
ACCENT = RGBColor(0xFF, 0xD7, 0x6E)   # золото
FONT   = 'Segoe UI'

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]

def slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s

def box(s, x, y, w, h, fill=None, line=None):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = fill
    else: sh.fill.background()
    if line: sh.line.color.rgb = line; sh.line.width = Pt(1.25)
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    return sh

def txt(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    """runs: список абзацев; абзац = список (text, size, color, bold)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(6)
        for (t, size, color, bold) in para:
            r = p.add_run(); r.text = t
            r.font.size = Pt(size); r.font.color.rgb = color
            r.font.bold = bold; r.font.name = FONT
    return tb

def title(s, text, sub=None):
    txt(s, 0.6, 0.35, 12.1, 1.0, [[(text, 30, TEXT, True)]])
    bar = box(s, 0.65, 1.15, 1.6, 0.06, fill=ACCENT)
    if sub:
        txt(s, 0.6, 1.3, 12.1, 0.5, [[(sub, 14, MUTED, False)]])

def card(s, x, y, w, h, head, body, head_color=ACCENT, body_size=13):
    box(s, x, y, w, h, fill=CARD, line=EDGE)
    txt(s, x + 0.25, y + 0.2, w - 0.5, 0.6, [[(head, 16, head_color, True)]])
    txt(s, x + 0.25, y + 0.75, w - 0.5, h - 0.95,
        [[(b, body_size, TEXT, False)] for b in body])

def stars(s, n=60, seed=7):
    rnd = seed
    for i in range(n):
        rnd = (rnd * 1103515245 + 12345) % (2**31)
        x = (rnd % 1280) / 100.0
        rnd = (rnd * 1103515245 + 12345) % (2**31)
        y = (rnd % 700) / 100.0
        d = 0.03 if i % 5 else 0.05
        sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0xAA, 0xB6, 0xFF)
        sh.line.fill.background(); sh.shadow.inherit = False

# ===== 1. Титул =====
s = slide(); stars(s)
txt(s, 1, 1.6, 11.3, 0.6, [[('ПРЕЗЕНТАЦИЯ ПРОЕКТА', 18, MUTED, False)]], PP_ALIGN.CENTER)
txt(s, 1, 2.3, 11.3, 1.6, [
    [('«АСТРАЛЬНЫЙ СКИТАЛЕЦ»', 44, TEXT, True)],
    [('Разработка коммерческой игры под Яндекс Игры', 22, ACCENT, False)],
], PP_ALIGN.CENTER)
txt(s, 1, 4.6, 11.3, 1.2, [
    [('2D top-down survival / rogue-lite · Vanilla JS + HTML5 Canvas', 15, MUTED, False)],
    [('Команда: ____________________', 16, TEXT, False)],
    [('Екатеринбург · 2026', 13, MUTED, False)],
], PP_ALIGN.CENTER)

# ===== 2. Идея и проблема =====
s = slide(); title(s, 'Идея проекта и проблема')
card(s, 0.6, 1.8, 6.0, 4.9, 'Проблема игрока', [
    'Казуальные браузерные игры часто не дают осмысленного выбора: игрок механически повторяет одно действие.',
    'Запрос аудитории — короткая сессия с настоящим решением «рискнуть или сыграть осторожно» и сохранением прогресса между устройствами.',
])
card(s, 6.85, 1.8, 6.0, 4.9, 'Проблема студента (учебный трек)', [
    'У студентов 1 курса нет опыта полного цикла веб-публикации: учебные проекты не доходят до реальной дистрибуции.',
    'Создание HTML5-игры решает комплексную задачу: интеграция SDK, сборка ≤100 МБ, облачные сохранения, реальная модерация.',
])
txt(s, 0.6, 6.85, 12.3, 0.5, [[('Идея: rogue-lite про пилота, который после крушения собирает энергию по секторам планеты и возвращает память — каждый контейнер это выбор «риск или награда».', 13, ACCENT, False)]])

# ===== 3. Целевая аудитория =====
s = slide(); title(s, 'Целевая аудитория', 'методика 5W Шеррингтона')
card(s, 0.6, 1.8, 6.0, 5.0, '01 / Игроки (B2C)', [
    'Профиль: мужчины и женщины 12–35 лет из РФ и СНГ, играют с ПК.',
    'Потребность: быстрый запуск сессии на 5–15 минут без скачивания.',
    'Мотивация: азарт «риск/награда», прогрессия, сюжет с выбором.',
    'Где: каталог и рекомендации Яндекс Игр.',
])
card(s, 6.85, 1.8, 6.0, 5.0, '02 / Создатели (Edu)', [
    'Профиль: студенты технических вузов, начинающие инди-разработчики.',
    'Потребность: полный цикл веб-разработки на открытом примере — SDK, оптимизация, модерация.',
    'Ценность: сильный кейс в портфолио, воспроизводимый код на GitHub.',
])

# ===== 4. Анализ конкурентов =====
s = slide(); title(s, 'Анализ конкурентов')
headers = ['Критерий сравнения', 'Казуальные игры на площадке', 'Курсы и видеогайды', 'Наш проект']
rows = [
    ['Реальная веб-публикация', 'Да', 'Редко (локально)', 'Да'],
    ['Выбор «риск / награда»', 'Редко', '—', 'Да: шансы видны до вскрытия'],
    ['Платформенная готовность', 'Да (закрытый код)', 'Нет (макеты)', 'Да (полное соответствие)'],
    ['Открытый код', 'Нет', 'Иногда', 'Да (учебный референс)'],
    ['Вес сборки', '5–100 МБ', '—', '≈0,1 МБ'],
]
tbl = s.shapes.add_table(len(rows)+1, 4, Inches(0.6), Inches(1.7), Inches(12.1), Inches(4.9)).table
tbl.columns[0].width = Inches(3.4); tbl.columns[1].width = Inches(3.3)
tbl.columns[2].width = Inches(2.6); tbl.columns[3].width = Inches(2.8)
def cell(c, t, bold=False, color=TEXT, size=13, fill=CARD):
    c.fill.solid(); c.fill.fore_color.rgb = fill
    c.margin_left = Inches(0.12); c.margin_top = Inches(0.06)
    tf = c.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = t
    r.font.bold = bold; r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = FONT
for j, h in enumerate(headers):
    cell(tbl.cell(0, j), h, bold=True, color=ACCENT, fill=RGBColor(0x23, 0x2D, 0x5C))
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        cell(tbl.cell(i+1, j), v, bold=(j == 3), color=(ACCENT if j == 3 else TEXT))

# ===== 5. Цели и задачи =====
s = slide(); title(s, 'Цели и задачи проекта')
box(s, 0.6, 1.65, 12.3, 1.05, fill=CARD, line=ACCENT)
txt(s, 0.85, 1.8, 11.9, 0.8, [[('SMART-цель: разработать и подготовить к публикации HTML5-игру с интеграцией Yandex SDK, пройти модерацию и получить публичную ссылку за 4 недели.', 15, TEXT, True)]])
steps = [
    ('01. Анализ', 'Требования платформы, ЦА (5W), аналоги, выбор стека.'),
    ('02. Разработка', 'Игровое ядро на Canvas: сектора, бой, экономика, сюжет.'),
    ('03. Интеграция', 'Yandex SDK: автопауза, облачные сейвы + localStorage.'),
    ('04. Релиз', 'Оптимизация ≤100 МБ, кросс-браузерность, модерация.'),
]
for i, (h, b) in enumerate(steps):
    card(s, 0.6 + i * 3.12, 3.0, 2.92, 3.4, h, [b], body_size=13)

# ===== 6. Уникальность =====
s = slide(); title(s, 'Уникальность проекта')
card(s, 0.6, 1.8, 6.0, 5.0, 'Геймплей осознанного риска', [
    '• Шанс охраны виден ДО вскрытия контейнера — решение за игроком.',
    '• Телеграфированные атаки всех врагов: видно замах — можно уклониться.',
    '• Rogue-lite смерть: теряется еда и часть лома, но не прогресс.',
    '• Сюжет с возвращением памяти и двумя концовками.',
])
card(s, 6.85, 1.8, 6.0, 5.0, 'Фокус на инфраструктуре', [
    '• Сборка ≈0,1 МБ — в 1000 раз меньше лимита платформы.',
    '• Ноль внешних зависимостей: вся графика — Canvas, весь звук — WebAudio.',
    '• Эталонная обработка pause/resume и облачных сохранений.',
    '• Открытый код: учебный референс полного цикла публикации.',
])

# ===== 7. Требования к продукту =====
s = slide(); title(s, 'Требования к продукту')
card(s, 0.6, 1.8, 3.94, 4.9, 'Технические', [
    'Экспорт в HTML5, лимит веса 100 МБ, кросс-браузерность (Chrome, Firefox, Safari, Edge), стабильная работа без внешних зависимостей.',
])
card(s, 4.79, 1.8, 3.94, 4.9, 'Интеграция SDK', [
    'Обязательная автопауза при сворачивании вкладки и асинхронные облачные сохранения savePlayerData / getPlayerData с резервом в localStorage.',
])
card(s, 8.98, 1.8, 3.94, 4.9, 'Платформенные', [
    'Русский интерфейс, соответствие регламентам платформы и контентной политике 12+, адаптация под 21:9 / 16:9 / 4:3.',
])

# ===== 8. Стек =====
s = slide(); title(s, 'Технологический стек')
card(s, 0.6, 1.8, 6.0, 5.0, 'Почему Vanilla JS?', [
    'Тяжёлые движки (Unity) генерируют громоздкие веб-билды, которые сложно вписать в лимиты без потери производительности.',
    'Нативный JS даёт минимальный вес, мгновенную загрузку и полный контроль над кодом — критично для модерации и учебных целей.',
])
card(s, 6.85, 1.8, 6.0, 5.0, 'Выбранные инструменты', [
    '• HTML5 Canvas API — лёгкий 2D-рендеринг (вся графика кодом).',
    '• Web Audio API — процедурный звук, ноль аудиофайлов.',
    '• Yandex Games SDK — авторизация, сейвы, события паузы.',
    '• Git + GitHub — версионный контроль и история итераций.',
    '• Node.js — 42 автотеста логики без браузера.',
    '• sdk-dev-proxy — локальная эмуляция платформы.',
])

# ===== 9. Схема работы =====
s = slide(); title(s, 'Схема работы приложения')
flow = [
    ('Инициализация', 'Загрузка /sdk.js → YaGames.init() → getPlayerData() (облако или localStorage) → LoadingAPI.ready().'),
    ('Игровой цикл', 'requestAnimationFrame: ввод → логика (голод, враги, бой) → рендер Canvas → HUD. Камера с зумом, частицы, телеграфы атак.'),
    ('Синхронизация', 'События игры меняют модель сейва → дебаунс 0,8 с → savePlayerData + localStorage. Автопауза по visibilitychange / blur / pagehide.'),
]
for i, (h, b) in enumerate(flow):
    card(s, 0.6 + i * 4.18, 1.9, 3.98, 4.4, h, [b])
    if i < 2:
        ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.48 + i * 4.18), Inches(3.85), Inches(0.4), Inches(0.5))
        ar.fill.solid(); ar.fill.fore_color.rgb = ACCENT; ar.line.fill.background(); ar.shadow.inherit = False

# ===== 10. Финальный геймплей =====
s = slide(); title(s, 'Финальный геймплей')
card(s, 0.6, 1.7, 12.3, 5.2, 'Что получилось', [
    '• 5 сюжетных секторов + бесконечные «дальние сектора» после второй концовки.',
    '• Контейнеры с видимым риском: лут сразу — или бой с охраной за награду ×1,5.',
    '• Бой: два оружия (Tab), криты 15 %, рывок с неуязвимостью, автоприцел; у всех врагов атаки с телеграфом.',
    '• Босс «Страж Ядра»: замах → удар по площади → окно уязвимости.',
    '• Экономика дефицита: голод, лом, магазин-синтезатор у шлюпки, баффы.',
    '• Облачные сохранения + автопауза: полное соответствие регламентам Яндекс Игр.',
    '• Сборка ≈100 КБ, 42 автотеста, открытый репозиторий на GitHub.',
], body_size=15)

# ===== 11. Финал =====
s = slide(); stars(s, seed=21)
txt(s, 1, 2.4, 11.3, 1.5, [
    [('Спасибо за внимание!', 40, TEXT, True)],
    [('Готовы ответить на ваши вопросы', 18, MUTED, False)],
], PP_ALIGN.CENTER)
txt(s, 1, 5.1, 11.3, 1.2, [
    [('Команда: ____________________', 16, TEXT, False)],
    [('«Астральный скиталец» · Разработка под Яндекс Игры · 2026', 13, ACCENT, False)],
], PP_ALIGN.CENTER)

prs.save('/tmp/Презентация_Астральный_скиталец.pptx')
print('pptx ok, слайдов:', len(prs.slides.__iter__.__self__._sldIdLst))
